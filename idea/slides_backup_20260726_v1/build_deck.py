#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the group-meeting deck (16:9) from the figures in figures/.
Image-first, few words, speaker notes embedded. Chinese font: Alibaba PuHuiTi 3.0."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(__file__)
FIG = os.path.join(HERE, "figures")
EMU_IN = 914400
CJK = "Alibaba PuHuiTi 3.0"

INK   = RGBColor(0x1F,0x2A,0x44); GRAY = RGBColor(0x6B,0x76,0x88)
BLUE  = RGBColor(0x2B,0x6C,0xB0); TEAL = RGBColor(0x2C,0x7A,0x7B)
RED   = RGBColor(0xC5,0x30,0x3A); AMBER= RGBColor(0xB7,0x79,0x1F)
VIO   = RGBColor(0x6B,0x46,0xC1); GREEN= RGBColor(0x2F,0x85,0x5A)
WHITE = RGBColor(0xFF,0xFF,0xFF); LGRAY= RGBColor(0xAE,0xB6,0xC2)
BG    = RGBColor(0xFF,0xFF,0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = 13.333, 7.5
BLANK = prs.slide_layouts[6]

def _ea(run, font=CJK):
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin","a:ea","a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", font)

def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    if isinstance(runs[0], tuple): runs = [runs]
    for i,para in enumerate(runs):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align
        for (s,size,color,bold) in para:
            r = p.add_run(); r.text = s
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            _ea(r)
    return tb

def rect(slide, x,y,w,h, fill, line=None, lw=1.0, rounded=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if rounded:
        try: shp.adjustments[0] = 0.06
        except Exception: pass
    return shp

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def footer(slide, n):
    txt(slide, 0.55, 7.06, 8, 0.3, [[("组会 · Do Not Mention / SCR / HalluSquatting → 我的 idea", 10, LGRAY, False)]])
    txt(slide, 12.0, 7.06, 0.9, 0.3, [[(str(n), 10, LGRAY, False)]], align=PP_ALIGN.RIGHT)

def bg_white(slide):
    r = rect(slide, -0.1,-0.1, SW+0.2, SH+0.2, BG);
    slide.shapes._spTree.remove(r._element); slide.shapes._spTree.insert(2, r._element)

def title_bar(slide, title, accent=TEAL, kicker=None):
    txt(slide, 0.6, 0.32, 11.8, 0.7, [[(title, 27, INK, True)]])
    rect(slide, 0.62, 1.02, 1.5, 0.055, accent)
    if kicker:
        txt(slide, 0.62, 0.05, 8, 0.3, [[(kicker, 12.5, accent, True)]])

def add_image_fit(slide, path, top=1.28, bottom=6.55, left=0.5, right=12.83):
    iw, ih = Image.open(path).size
    aw, ah = (right-left), (bottom-top)
    scale = min(aw/(iw/ EMU_IN*EMU_IN)*1.0, 1)  # placeholder
    # compute by inches assuming 200dpi-ish irrelevant; fit by aspect
    ar = iw/ih
    if aw/ah > ar:   # height-bound
        h = ah; w = ah*ar
    else:
        w = aw; h = aw/ar
    x = left + (aw-w)/2; y = top + (ah-h)/2
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))

def takeaway(slide, text, color=TEAL):
    y=6.62; h=0.6
    rect(slide, 0.5, y, 12.33, h, RGBColor(0xF2,0xF5,0xFA), rounded=True)
    rect(slide, 0.5, y, 0.13, h, color)
    txt(slide, 0.78, y, 12.0, h, [[("一句话：", 15, color, True),(text, 15, INK, False)]],
        anchor=MSO_ANCHOR.MIDDLE)

# ---- slide builders -------------------------------------------------------
def s_title():
    s = prs.slides.add_slide(BLANK); bg_white(s)
    rect(s, 0, 0, 0.35, SH, TEAL)
    rect(s, 0, SH-0.35, SW, 0.35, INK)
    txt(s, 1.0, 1.5, 11.5, 1.2, [[("组会论文分享", 40, INK, True)]])
    txt(s, 1.0, 2.7, 11.6, 1.4, [
        [("主讲  ", 24, TEAL, True), ("“Do Not Mention This to the User”", 24, INK, True)],
        [("       Detecting & Understanding Malicious Agent Skills in the Wild · USENIX Security 2026", 15, GRAY, False)],
        [("补充  ", 24, TEAL, True), ("SCR（组合风险）  ｜  HalluSquatting（幻觉抢注）", 22, INK, True)],
    ])
    rect(s, 1.02, 4.7, 11.2, 1.15, RGBColor(0xE7,0xDF,0xF7), rounded=True)
    rect(s, 1.02, 4.7, 0.14, 1.15, VIO)
    txt(s, 1.35, 4.7, 10.7, 1.15, [
        [("我的方向：", 18, VIO, True), ("当用户只给普通任务时，agent 会不会自主“获取”一个尚未安装的恶意 skill？", 18, INK, True)],
        [("三篇论文正好拼出「恶意供给 → 促成安装 → 真实执行」，我来测这条链能不能自主接通。", 15, GRAY, False)],
    ])
    txt(s, 1.0, 6.35, 11, 0.4, [[("汇报人：（你的名字）    2026-07", 13, LGRAY, False)]])
    notes(s, "开场30秒：今天分享三篇 agent skill 安全的论文，主讲 Do Not Mention，"
             "它是 USENIX Security 2026 已录用的大规模测量；再用 SCR 和 HalluSquatting 两篇补充。"
             "更重要的是，我会把它们和我正在做的 idea 串起来——一句话：当用户只给普通任务时，"
             "agent 会不会自己去发现、安装、调用一个恶意 skill。全程约 30 分钟。")

def s_divider(kicker, title, sub, color):
    s = prs.slides.add_slide(BLANK)
    rect(s, -0.1,-0.1, SW+0.2, SH+0.2, color)
    txt(s, 1.0, 2.6, 11.3, 0.6, [[(kicker, 20, RGBColor(0xFF,0xFF,0xFF), True)]])
    txt(s, 1.0, 3.15, 11.3, 1.3, [[(title, 40, WHITE, True)]])
    txt(s, 1.02, 4.5, 11.3, 1.0, [[(sub, 20, RGBColor(0xE8,0xEE,0xF6), False)]])
    rect(s, 1.04, 4.35, 2.2, 0.05, RGBColor(0xFF,0xFF,0xFF))
    return s

def s_img(title, img, tk=None, tkcolor=TEAL, accent=TEAL, kicker=None, note="", n=0,
          top=1.3, bottom=(6.5)):
    s = prs.slides.add_slide(BLANK); bg_white(s)
    title_bar(s, title, accent, kicker)
    b = 6.5 if tk is None else 6.45
    add_image_fit(s, os.path.join(FIG, img), top=top, bottom=b)
    if tk: takeaway(s, tk, tkcolor)
    if note: notes(s, note)
    footer(s, n)
    return s

def s_text(title, blocks, accent=TEAL, kicker=None, note="", n=0):
    s = prs.slides.add_slide(BLANK); bg_white(s)
    title_bar(s, title, accent, kicker)
    txt(s, 1.1, 1.7, 11.1, 4.8, blocks, anchor=MSO_ANCHOR.TOP)
    if note: notes(s, note)
    footer(s, n)
    return s

# ============================ BUILD ========================================
n=1
s_title();
# 2 roadmap
s_img("今天的主线：一个问题的三段", "fig_roadmap.png",
      tk="恶意供给 → 促成安装 → 真实执行；我落在“普通任务能否自主获取”这一段。",
      tkcolor=VIO, accent=VIO, kicker="ROADMAP",
      note="这张图是今天的骨架。真实的 agent 供应链风险要同时成立三件事：①市场里真有恶意 skill（Do Not Mention）；"
           "②就算每个 skill 单独干净，组合或背书也能促成有害安装（SCR）；③连装哪个都能被 AI 幻觉劫持、一路打到远程代码执行（HalluSquatting）。"
           "但三篇都没回答最下面那条紫色的问题：用户只给普通任务、目标一开始没装时，agent 会不会自己走完发现→安装→调用。这就是我的落点。", n=2); n=3
# 3 background
s_img("先铺个底：什么是 agent skill？", "fig_skill_bg.png", accent=TEAL, kicker="背景",
      note="给没接触过的同学 30 秒背景。Agent 就是会自己规划、会调用工具（开终端、读写文件、联网）的 LLM 应用。"
           "Skill 是第三方插件，一个文件夹，核心是 SKILL.md：一半是给人看的自然语言说明，一半是真正执行的脚本。"
           "用户一句“帮我装个天气 skill”，agent 就去公开市场下载安装。危险有三点：默认信任、本地权限、以及——自然语言本身就是攻击面，坏话可以直接写在说明里。", n=3); n=4

# ===== Do Not Mention =====
s_divider("主讲 · MAIN", "“Do Not Mention This to the User”",
          "Detecting & Understanding Malicious Agent Skills in the Wild · USENIX Security 2026 · Yi Liu et al.", BLUE); n=5
s_img("一个真实的恶意 skill 长什么样", "dnm_fig1_shadow.png", accent=BLUE, kicker="Do Not Mention",
      tk="说明写“安全的计算器”，脚本里却藏了反弹 shell —— 文档说的 ≠ 实际做的。", tkcolor=BLUE,
      note="先看一个论文里的真实样本，这也是标题的由来。左边 SKILL.md 写着 math-calculator、“安全地做算术”；"
           "右边 calculate.py 第 32 行却是一段反弹 shell：连到攻击者的 20.120.229.246:4444，把终端交出去。"
           "文档和行为完全不一致——论文把这种“文档掩盖真实行为”叫 shadow feature，标题“别跟用户提这个”就是这类 skill 里反复出现的指令。", n=5); n=6
s_img("它怎么测：从 9.8 万到 157", "dnm_fig3_funnel.png", accent=BLUE, kicker="测量漏斗",
      tk="0.16% 是“行为确认恶意”的高精度下界，不是“生态只有 0.16% 有风险”。", tkcolor=BLUE,
      note="这是全篇骨架。两个公开市场共 98,380 个 skill；静态规则筛出 4,287 个可疑（4.4%）；"
           "放进沙箱动态执行，762 个真的触发了可疑运行时行为；两位研究者独立人工复核，确认 157 个是真恶意，标注出 632 个漏洞、13 种技术。"
           "每一层都在收紧证据强度。注意 157/98,380 = 0.16% 是“高精度行为确认”的下界，因为很多可疑样本因沙箱时限或休眠触发没被激活；"
           "但也不能反过来说真实比例一定高很多。顺便点一个坑：SCR 引用的“26.1% 有漏洞”不是这篇，是另一篇的静态筛查口径，别混。", n=6); n=7
s_img("为什么可信：四道关", "fig_dnm_verify.png", accent=BLUE, kicker="验证质量",
      note="它的可信度来自方法。Docker 隔离沙箱，全程监控网络、系统调用、文件；埋 honeypot 诱饵凭证看它会不会去偷；"
           "用三种方式触发（元数据、LLM 生成输入、多轮调用）；而且只有观察到真实坏行为才进人工双复核，一致性 κ=0.91。"
           "结果是行为验证精度 99.6%，比最强的纯静态基线（≤1.1%）高约 90 倍。责任披露后 157 个全部下架。", n=7); n=8
s_img("三个发现：RQ1 / RQ2 / RQ3", "fig_dnm_rq.png", accent=BLUE, kicker="核心结论",
      tk="重点是后半段的分析：攻击是蓄意分层、分工明确、且逃避随复杂度升级。", tkcolor=BLUE,
      note="这篇不是“扫完文件就结束”，后半段的分析才是贡献。RQ1 全貌：632 漏洞、13 技术、6 阶段，平均每个恶意 skill 含 4.03 个漏洞——是蓄意分层；84.2% 的漏洞就写在 SKILL.md 的自然语言里。"
           "RQ2 策略：用统计检验发现两种“负相关”的攻击原型——Data Thieves 专做数据外泄，Agent Hijackers 专做指令劫持，攻击者二选一专精；而且单一“工业化”行为者用模板化品牌仿冒贡献了 54.1%。"
           "RQ3 逃避：越高级的攻击越隐蔽，shadow features 从最低级的 0% 升到最高级的 100%，还会武器化平台自带的信任机制。", n=8); n=9
s_text("它回答了什么，没回答什么", [
    [("✓  回答了：", 24, GREEN, True)],
    [("     外部市场里真有行为确认的恶意 skill —— 而且已经规模化、隐蔽化供应。", 22, INK, False)],
    [("", 10, INK, False)],
    [("✗  没回答：", 24, RED, True)],
    [("     当用户只给普通能力任务时，agent 会不会“自主”把这些第三方 skill 发现、安装、调用进来？", 22, INK, False)],
    [("", 10, INK, False)],
    [("→  ", 22, VIO, True), ("生态测量给了“现实性锚点”，但替代不了 acquisition 的端到端实验。这正是我的切入点。", 22, VIO, True)],
], accent=BLUE, kicker="承上启下",
   note="小结主讲、接到我的问题。它回答了“坏 skill 真实存在吗”——有，而且规模化、隐蔽化。"
        "但它没回答“用户只给普通任务时，agent 会不会自己把它获取进来”。所以它给了现实性，但不能替代获取链路的端到端实验。带着这个缺口看两篇补充。", n=9); n=10

# ===== SCR =====
s_divider("补充① · SUPPLEMENT", "SCR：Benign in Isolation, Harmful in Composition",
          "Skill Composition Risk · arXiv 2606.15242 · 华东师大 / A*STAR（预印本）", TEAL); n=11
s_img("核心：单个都干净，组合起来出事", "scr_fig1_chain.png", accent=TEAL, kicker="SCR",
      tk="三个 skill 单独审都合规，串在一起——“清单+已审核+建议”被当成授权，就越权开放了资源。", tkcolor=TEAL,
      note="类比：保安交一张“这些门没锁”的清单，主管说“任务挺正常”，门禁员照着就把门开了——没人单独做错，门却开了。"
           "SCR 把“人”换成“skill”。论文这张图：文件审计 skill 输出目标清单，安全审查 skill 输出“建议批准”，访问管理 skill 消费这两个信号就改了权限。"
           "危险不在任何一个 skill，而在它们之间被激活的“组合路径”。所以它主张：审核单位要从“单个 skill”升级到“激活的路径”，而且只看真实的下游状态改变，不看嘴上说。", n=11); n=12
s_img("结果：单独 vs 组合，三种机制", "fig_scr_results.png", accent=TEAL, kicker="SCR 结果",
      tk="TrustLift 最猛：一个“已审核”背书，把有害安装从 1.1% 抬到 83.9%。", tkcolor=TEAL,
      note="它做了 SCR-Bench，三个子基准，都在沙箱里以“真实状态改变”计分，跑了约十个大模型。"
           "能力流 CapFlow：单独几乎 0%，组合后平均 33.6%，DeepSeek-V4 超过 90%。"
           "信任转移 TrustLift 最猛：没有背书时有害安装平均 1.1%，加一个“安全审查已通过”的背书就冲到 83.9%，五个后端里四个超过 96.5%。"
           "授权混淆 AuthBlur：把“建议”当“授权”，风险从 15.7% 升到 27%。一句话：单独看无害，组合或背书后风险跃升。"
           "注意第一列 Neutral 用的是中性任务语言——所以我不能说自己“第一个测自主/中性触发”。", n=12); n=13
s_text("SCR 的边界 → 接我的问题", [
    [("SCR 已经很接近，但它：", 22, INK, True)],
    [("   • 全部在 mock 沙箱、观测“模拟状态改变”，没有真实市场 / 真实落盘 / 注册；", 21, INK, False)],
    [("   • TrustLift 从一个“已经出现的下游安装请求”起步，不是普通任务自发去发现；", 21, INK, False)],
    [("", 8, INK, False)],
    [("所以我不能再说：", 20, RED, True),("“第一个研究 install”“第一个测自主触发”“SCR 只测预装 skill”。", 20, INK, False)],
    [("", 8, INK, False)],
    [("→  ", 22, VIO, True),("可守的区别：普通任务触发 + 真实落盘/注册/调用 的“同一条轨迹”逐段测量。", 22, VIO, True)],
], accent=TEAL, kicker="边界",
   note="必须诚实说清 SCR 的边界，也顺便修正我以前的过度主张。SCR 用 mock 沙箱是为了拿到干净的路径级 ground truth，代价是没覆盖真实市场和真实落盘；"
        "而且 TrustLift 是从一个“已经出现的安装请求”开始的。所以我不能再说“第一个研究 install”或“第一个测自主触发”，这些都被它占了。"
        "我可守的区别是：从普通能力任务触发、目标未装、走真实落盘/注册/调用的同一条轨迹逐段测。", n=13); n=14

# ===== HalluSquatting =====
s_divider("补充② · SUPPLEMENT", "HalluSquatting：Beware of Agentic Botnets",
          "抢注 AI 的“幻觉” · arXiv 2607.07433 · Tel Aviv U / Technion / Intuit（预印本）", RED); n=15
s_img("机制：抢注“AI 记错的名字”", "hallu_fig1_threat.png", accent=RED, kicker="HalluSquatting",
      tk="门槛极低：坏人只要能看趋势 + 在公开平台注册一个名字，等 AI 自己“记错”送上门。", tkcolor=RED,
      note="类比 typosquatting，但抢的不是人的错别字，是 AI 的错。你让 agent “clone 那个 librepods”“装个 skill vetter”，"
           "AI 会自信地编出一个错误的地址（错的 GitHub owner，或错的 slug）。坏人提前统计 AI 最常编哪个错名、抢先注册、塞入恶意提示词。"
           "看这张威胁模型图：用户触发→AI 规划→幻觉出资源名→框架去取回→上下文投毒→工具调用，最后装 bot。因为幻觉跨模型可预测，占坑一次能命中一大片，甚至组 botnet。"
           "威胁模型很弱：坏人只要能看趋势榜、能注册一个资源，不需要攻破用户。", n=15); n=16
s_img("证据①：AI 到底多爱记错", "fig_hallu_halluc.png", accent=RED, kicker="证据 1/3",
      tk="只要资源比模型训练时间新，AI 几乎必然编错 owner —— 而且编得有规律、可抢注。", tkcolor=RED,
      note="第一组证据，一万四千次查询。2025 年新出的热门仓库，平均 92.4% 编错 owner，很多组合是 100% 全错；"
           "2013–2018 的老仓库作为对照只有 0.9%。差别就在训练数据：老的见过、记得住，新的只能编。skill 安装场景同理，可高达 100%。"
           "关键不是“偶尔手滑”，而是“系统性、可预测”，坏人正是统计这个分布去占坑。", n=16); n=17
s_img("证据②：skill 市场的两个结构缺口", "fig_hallu_v1v2.png", accent=RED, kicker="证据 2/3",
      tk="不是用户打错字，是市场命名结构 + AI 解析习惯，让“记错”成了默认结果。", tkcolor=RED,
      note="第二组证据，专门针对 skill 安装。V1 掉词：用户说 install skill-vetter（市场第二热门），"
           "因为引导文里“skill”这个词太多，模型把它当噪声删掉，解析成 slug vetter，80–100% 命中——坏人注册 vetter 就行。"
           "V2 名不符实：市场把显示名和真实 slug 分开索引，用户复制显示名去装，模型只能编一个，100% 落到可占坑的名字。"
           "最惨是非英文文档的 skill，英文搜索直接把正牌挤出前十，连纠错路径都断了。", n=17); n=18
s_img("证据③：真的打到 RCE", "fig_hallu_e2e.png", accent=RED, kicker="证据 3/3",
      tk="真实产品上端到端命中：代码助手 20–65%，个人助手 40–100%；装进去后聊天外泄 96%。", tkcolor=RED,
      note="第三组证据，真实产品上的端到端。代码助手（Cursor、Windsurf、Cline 等六个）clone 场景 20–65% 打到工具调用或 RCE；"
           "个人助手（OpenClaw、ZeroClaw、NanoClaw）install 场景 40–100%。装进去之后，聊天记录外泄 96%、终端反弹 shell 84%。"
           "要讲清楚它是受控分阶段：公开平台只注册良性替身测占坑，真实 RCE 载荷是在研究者自己机器上测的，不是公开投毒。"
           "一个细节：只有会“运行前完整检查脚本”的后端偶尔拒绝，几乎所有 Claude 后端都默默中招。", n=18); n=19
s_text("HalluSquatting 的边界 → 接我的问题", [
    [("它已经打到真实 RCE（40–100%），所以我不能说“第一个测真实执行”。但：", 22, INK, True)],
    [("", 8, INK, False)],
    [("   它的触发是 ", 21, INK, False),("用户明确要求 clone / install", 21, RED, True),("（用户本来就有获取意图，只是被幻觉+抢注劫持）。", 21, INK, False)],
    [("", 8, INK, False)],
    [("→  ", 22, VIO, True),("我要测的是：用户", 22, VIO, True),("完全没有获取意图", 22, RED, True),
     ("、只给普通能力任务时，agent 会不会自发把 skill 装进来。", 22, VIO, True)],
    [("", 8, INK, False)],
    [("      两者威胁模型不同：一个是“被劫持的显式安装”，一个是“任务诱导的自主安装”。", 19, GRAY, False)],
], accent=RED, kicker="边界",
   note="同样诚实收边界。它在真实产品上打到 RCE，所以“第一个测真实执行”也被占了。"
        "但它的触发是用户明确要 clone/install——用户已经有获取意图，只是被幻觉劫持到恶意资源。"
        "我测的是用户零获取意图、只给普通任务时，agent 会不会自发去装。威胁模型不同：被劫持的显式安装 vs 任务诱导的自主安装。", n=19); n=20

# ===== My idea =====
s_divider("我的 idea · MY WORK", "三篇 → 我要测的“同一条轨迹”",
          "Task-induced Autonomous Acquisition：普通任务能否自主接通 发现→安装→调用", VIO); n=21
s_img("三篇各覆盖一段，我测同一条链", "fig_integration.png", accent=VIO, kicker="整合",
      tk="三篇分别做了“供给 / 促装 / 执行”；没人把“普通任务→自主全链→逐段归因”放进同一条轨迹。", tkcolor=VIO,
      note="这是最关键的一张图。把 agent 获取一个 skill 的完整生命周期横着摆开：能力缺口→外部发现→发出安装→解析执行→落盘→注册→调用→载荷。"
           "Do Not Mention 在最左边——市场里有恶意供给在等着。SCR 覆盖“发出→执行”，但在 mock 环境、从已有安装请求起步。"
           "HalluSquatting 覆盖“显式 install→落盘→调用→RCE”，但用户已有获取意图。"
           "没有一篇把最下面那条紫色的路径——普通任务触发、目标未装、agent 自主走完整条链、并逐段测断点——放进同一条轨迹。那个问号就是我怀疑最容易断的地方。", n=21); n=22
s_img("我的框架：把风险拆成逐段存活率", "fig_my_funnel.png", accent=VIO, kicker="方法",
      tk="不报一个总 ASR，而是报每一段的条件存活率，指出真正的瓶颈在哪。", tkcolor=VIO,
      note="方法上，我不报一个笼统的攻击成功率，而是把端到端风险写成逐段条件概率的乘积，测每一段的存活率，找真正的瓶颈。"
           "紫色高亮的“发出→执行”是我最怀疑会断的地方——先导调试里模型已经吐出了 install 调用，却因为 tool budget 没真执行。"
           "必须强调：这张图是示意，每段的真实数字要由 pilot 测出来，不是既有结论。", n=22); n=23
s_img("我已有的先导抓手", "fig_pilot.png", accent=VIO, kicker="先导结果",
      tk="两个先导现象都要带条件讲：它们是 idea 演化的证据，不是普适结论。", tkcolor=VIO,
      note="我已经有两个先导抓手。左边：预装条件下，恶意 skill 的选中率 Vanilla 0/20、粗蹭热度 2/20、5 轮优化 1/20——只在单模型/单框架/20 任务下成立，是先导观察不是定律。"
           "右边就是那个测量陷阱：模型吐出了 install_skill，但默认 3 轮 tool budget 被耗尽，安装从未真正执行。"
           "它说明 emitted / parsed / executed 必须拆开测，也正是我下一步要修的 instrumentation。", n=23); n=24
s_text("讨论 & 下一步", [
    [("下一步（先做小 pilot，再决定要不要扩大）", 22, VIO, True)],
    [("   1. 修 instrumentation：可配 tool budget，拆开 emitted / parsed / executed / on-disk / registered / invoked；", 19, INK, False)],
    [("   2. 真实落盘 + manifest + registry 三重验证“真的装上了”；", 19, INK, False)],
    [("   3. 60–100 次小 pilot，按预先规则 Go / No-Go —— 只有出现稳定、可归因、可复现的结构才继续。", 19, INK, False)],
    [("", 10, INK, False)],
    [("想请大家讨论", 22, TEAL, True)],
    [("   Q1  一个“全链逐段测量”要出现什么级别的结构性结果，才够成一篇 measurement paper？", 19, INK, False)],
    [("   Q2  要不做成玩具市场，正式实验至少需要哪些真实要素：真实 registry、真实落盘、多个 scaffold？", 19, INK, False)],
], accent=VIO, kicker="收尾",
   note="收尾讲下一步和讨论题。下一步很克制：先修好逐段的 instrumentation，把安装做成真实落盘+manifest+registry 三重验证，"
        "跑 60–100 次小 pilot，按预先写好的规则做 Go/No-Go——只有出现稳定、可归因、可复现的结构才继续，否则就停，不硬凑故事。"
        "两个讨论题：一是全链测量要多强的结构才够发论文；二是要多真实的系统要素才不算玩具。谢谢，请老师和大家指点。", n=24); n=25
# backup: 口径红线
s_text("（备用）汇报口径红线", [
    [("被追问时，守住这几条：", 20, INK, True)],
    [("   • 不说“第一个测 install / E2E / 自主触发” —— SCR、HalluSquatting 已分别覆盖。", 18, INK, False)],
    [("   • “26.1% 有漏洞”是另一篇的静态筛查口径，被 SCR 误引；本篇是 0.16% 行为确认。", 18, INK, False)],
    [("   • 我的 0/2/1 是单模型先导观察，不是“明确任务天然安全”的普适结论。", 18, INK, False)],
    [("   • 三篇百分比不能直接横比：起点、终点、成功定义都不同。", 18, INK, False)],
    [("   • research gap ≠ novelty：贡献要来自新测量 / 新现象 / 方法 / 防御，低装机率也要有结构才有价值。", 18, INK, False)],
], accent=GRAY, kicker="BACKUP",
   note="这页是备用，被追问时用。关键红线：不抢“第一个”；26.1% 别错当本篇；自己的先导结果带条件；三篇不横比；gap 不等于贡献。", n=25)

out = os.path.join(HERE, "组会_Do_Not_Mention_三篇.pptx")
prs.save(out)
print("saved", out, "| slides:", len(prs.slides._sldIdLst))
